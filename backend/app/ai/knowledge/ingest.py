"""Document parsing and deterministic chunking for the RAG ingestion pipeline."""
from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass


@dataclass(frozen=True)
class ParsedDocument:
    file_name: str
    content_type: str
    text: str
    pages: tuple[str, ...] = ()


def content_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def parse_document(file_name: str, data: bytes, content_type: str | None = None) -> ParsedDocument:
    """Parse supported formats without trusting a client-provided MIME type."""
    suffix = file_name.lower().rsplit(".", 1)[-1] if "." in file_name else ""
    if suffix in {"txt", "md", "markdown"}:
        text = data.decode("utf-8-sig", errors="replace")
        return ParsedDocument(file_name, content_type or "text/plain", text)
    if suffix == "pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ValueError("PDF 解析依赖未安装，请安装 pypdf") from exc
        import io
        pages = tuple((page.extract_text() or "") for page in PdfReader(io.BytesIO(data)).pages)
        return ParsedDocument(file_name, content_type or "application/pdf", "\n\n".join(pages), pages)
    if suffix == "docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise ValueError("DOCX 解析依赖未安装，请安装 python-docx") from exc
        import io
        paragraphs = tuple(p.text for p in Document(io.BytesIO(data)).paragraphs if p.text.strip())
        return ParsedDocument(file_name, content_type or "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "\n".join(paragraphs))
    raise ValueError("仅支持 TXT、Markdown、PDF 和 DOCX 文件")


def chunk_text(text: str, *, max_chars: int = 1200, overlap: int = 120) -> list[tuple[str, str]]:
    """Split on headings/paragraphs first, returning (content, locator)."""
    if max_chars <= 0 or overlap < 0 or overlap >= max_chars:
        raise ValueError("chunk 参数无效")
    blocks = [b.strip() for b in re.split(r"\n\s*\n+", text.replace("\r\n", "\n")) if b.strip()]
    result: list[tuple[str, str]] = []
    current_heading = ""
    for block in blocks:
        heading = next((line.strip() for line in block.splitlines() if line.lstrip().startswith("#")), "")
        if heading:
            current_heading = heading
        else:
            heading = current_heading
        start = 0
        while start < len(block):
            end = min(len(block), start + max_chars)
            piece = block[start:end].strip()
            if piece:
                result.append((piece, heading))
            if end == len(block):
                break
            start = end - overlap
    return result
