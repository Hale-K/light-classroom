"""排课完整包 Excel：说明 + 教师课时关系 + 教师课时 + 教师分课表 + 各班课表。"""
from __future__ import annotations

import io
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SUBJECT_SHORT = {
    "语文": "语",
    "数学": "数",
    "英语": "英",
    "物理": "物",
    "化学": "化",
    "生物": "生",
    "政治": "政",
    "历史": "历",
    "地理": "地",
    "体育": "体",
    "音乐": "音",
    "美术": "美",
    "心理": "心",
    "自主学习": "自",
    "班主任晚课": "班",
}

SUBJECT_ORDER = [
    "语文", "数学", "英语", "物理", "化学", "生物", "政治", "历史", "地理",
    "体育", "音乐", "美术", "心理",
]

SUBJECT_HEADER_FILL = {
    "语文": "5B9BD5",
    "数学": "ED7D31",
    "英语": "70AD47",
    "物理": "7030A0",
    "化学": "00B0F0",
    "生物": "00B050",
    "政治": "C00000",
    "历史": "FFC000",
    "地理": "4472C4",
    "体育": "FF6B6B",
    "音乐": "9B59B6",
    "美术": "E67E22",
    "心理": "1ABC9C",
}

PACK_SHEET_KEYS = (
    "cover",
    "teacher_relation",
    "teacher_hours",
    "teacher_grid",
    "class_timetables",
)

PACK_SHEET_LABELS = {
    "cover": "说明",
    "teacher_relation": "教师课时关系",
    "teacher_hours": "教师课时",
    "teacher_grid": "教师分课表",
    "class_timetables": "各班课表",
}

DAY_HEADERS = ["一", "二", "三", "四", "五", "单六", "双六"]


@dataclass
class PackSlot:
    class_id: int
    class_name: str
    subject_id: int
    subject_name: str
    teacher_id: int | None
    teacher_name: str
    weekday: int
    period: int
    week_parity: str


def _wt(parity: str) -> float:
    return 1.0 if parity == "all" else 0.5


def _short_subject(name: str) -> str:
    if name == "自主学习":
        return ""
    if name in SUBJECT_SHORT:
        return SUBJECT_SHORT[name]
    return (name or "?")[:1]


def _cell_label(entries: list[tuple[str, str]]) -> str:
    entries = [(s, p) for s, p in entries if s]
    if not entries:
        return ""
    alls = [s for s, p in entries if p == "all"]
    odds = [s for s, p in entries if p == "odd"]
    evens = [s for s, p in entries if p == "even"]
    if alls and not odds and not evens:
        return alls[0]
    parts: list[str] = []
    if odds:
        parts.append(odds[0])
    if evens:
        parts.append(evens[0])
    if alls and not parts:
        parts.append(alls[0])
    elif alls and parts:
        parts = alls + parts
    return "/".join(dict.fromkeys(parts))


def _class_no(name: str) -> str:
    m = re.search(r"(\d+)", name.replace("（", "(").replace("）", ")"))
    return m.group(1) if m else name


def _sheet_title(name: str) -> str:
    return name.replace("（", "(").replace("）", ")")[:31]


def plan_evening_hours(odd: int | None, even: int | None, parity: str | None) -> float:
    o = int(odd or 0)
    e = int(even or 0)
    p = (parity or "all").lower()
    if p in {"none", "无"}:
        return 0.0
    if p == "odd":
        return 1.0 if o > 0 else 0.0
    if p == "even":
        return 1.0 if e > 0 else 0.0
    if o > 0 and e > 0:
        return 1.0
    if o > 0 or e > 0:
        return 0.5
    return 0.0


def _safe_sheet_name(wb: Workbook, title: str) -> str:
    safe = title[:31] or "Sheet"
    base, n = safe, 1
    while safe in wb.sheetnames:
        suffix = f"_{n}"
        safe = base[: 31 - len(suffix)] + suffix
        n += 1
    return safe


def normalize_pack_sheets(sheets: list[str] | None) -> set[str]:
    """None = 默认全选；空列表 = 未勾选。"""
    if sheets is None:
        return set(PACK_SHEET_KEYS)
    return {key for key in sheets if key in PACK_SHEET_LABELS}


def pack_title_and_filename(grade_names: list[str], stamp: str | None = None) -> tuple[str, str]:
    stamp = stamp or datetime.now().strftime("%Y%m%d_%H%M%S")
    uniq = sorted({name for name in grade_names if name})
    if len(uniq) == 1:
        prefix = uniq[0].replace("年级", "")
        title = f"{prefix}-排课完整包"
    elif uniq:
        title = "排课完整包"
        prefix = "多年级"
    else:
        title = "排课完整包"
        prefix = "排课"
    return title, f"{prefix}-排课完整包_{stamp}.xlsx"


def build_timetable_pack_xlsx(
    *,
    pack_title: str,
    academic_year: str,
    term: str,
    scope_label: str,
    evening_start: int,
    slots: list[PackSlot],
    class_rows: list[tuple[int, str]],
    plan_by_cs: dict[tuple[int, int], dict[str, float]],
    teaching_assignments: list[dict[str, Any]],
    activity_subject_ids: set[int],
    include_sheets: set[str] | list[str] | None = None,
) -> bytes:
    include = normalize_pack_sheets(list(include_sheets) if include_sheets is not None else None)
    if not include:
        raise ValueError("请至少勾选一种工作表")

    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    title_fill = PatternFill("solid", fgColor="FFF2CC")
    head_fill = PatternFill("solid", fgColor="D9E2F3")
    eve_fill = PatternFill("solid", fgColor="FFF2CC")
    sum_head = PatternFill("solid", fgColor="4472C4")
    sum_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
    pe_font = Font(name="微软雅黑", color="C00000", size=11)
    normal = Font(name="微软雅黑", size=11)
    bold = Font(name="微软雅黑", bold=True, size=11)
    cell_font = Font(name="微软雅黑", size=10)
    white_bold = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    diff_fill = PatternFill("solid", fgColor="FCE4D6")

    wb = Workbook()
    wb.remove(wb.active)

    class_names = {cid: name for cid, name in class_rows}
    teacher_names: dict[int, str] = {}
    subject_names: dict[int, str] = {}
    for slot in slots:
        class_names.setdefault(slot.class_id, slot.class_name)
        if slot.teacher_id is not None:
            teacher_names[slot.teacher_id] = slot.teacher_name or str(slot.teacher_id)
        subject_names[slot.subject_id] = slot.subject_name

    need_teacher = bool(include & {"teacher_relation", "teacher_hours", "teacher_grid"})
    buckets: dict[tuple[int, int, int], dict[str, float]] = defaultdict(
        lambda: {"weekday": 0.0, "saturday": 0.0, "evening": 0.0}
    )
    plan_rows: dict[tuple[int, int, int], dict[str, float]] = {}
    if need_teacher:
        for slot in slots:
            if slot.teacher_id is None or slot.subject_id in activity_subject_ids:
                continue
            key = (slot.teacher_id, slot.subject_id, slot.class_id)
            w = _wt(slot.week_parity)
            if slot.period >= evening_start:
                buckets[key]["evening"] += w
            elif slot.weekday == 6:
                buckets[key]["saturday"] += w
            else:
                buckets[key]["weekday"] += w
        for row in teaching_assignments:
            tid = row.get("teacher_id")
            if tid is None:
                continue
            tid_i, sid, cid = int(tid), int(row["subject_id"]), int(row["class_id"])
            teacher_names.setdefault(tid_i, str(row.get("teacher_name") or tid_i))
            subject_names.setdefault(sid, str(row.get("subject_name") or sid))
            plan_rows[(tid_i, sid, cid)] = dict(
                plan_by_cs.get((cid, sid), {"weekday": 0.0, "saturday": 0.0, "evening": 0.0})
            )

    if "cover" in include:
        cover = wb.create_sheet("说明", 0)
        cover["A1"] = pack_title
        cover["A1"].font = Font(name="微软雅黑", bold=True, size=16)
        cover["A2"] = f"学年学期：{academic_year} 第{term}学期"
        cover["A3"] = f"导出范围：{scope_label}"
        cover["A4"] = f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        cover["A6"] = "本次包含工作表"
        line = 7
        for key in PACK_SHEET_KEYS:
            if key in include:
                cover[f"A{line}"] = f"· {PACK_SHEET_LABELS[key]}"
                line += 1
        cover.column_dimensions["A"].width = 72

    if "teacher_relation" in include:
        rel = wb.create_sheet("教师课时关系")
        headers = [
            "教师", "学科", "班级",
            "白天实排", "白天方案", "白天差",
            "周六实排", "周六方案", "周六差",
            "晚课实排", "晚课方案", "晚课差",
        ]
        for c, h in enumerate(headers, 1):
            cell = rel.cell(row=1, column=c, value=h)
            cell.fill = sum_head
            cell.font = sum_font
            cell.alignment = center
            cell.border = thin
        keys = sorted(
            set(buckets) | set(plan_rows),
            key=lambda k: (
                teacher_names.get(k[0], ""),
                SUBJECT_ORDER.index(subject_names.get(k[1], ""))
                if subject_names.get(k[1], "") in SUBJECT_ORDER
                else 999,
                class_names.get(k[2], ""),
            ),
        )
        for i, key in enumerate(keys, start=2):
            tid, sid, cid = key
            actual = buckets.get(key, {"weekday": 0.0, "saturday": 0.0, "evening": 0.0})
            plan = plan_rows.get(key, {"weekday": 0.0, "saturday": 0.0, "evening": 0.0})
            vals = [
                teacher_names.get(tid, tid),
                subject_names.get(sid, sid),
                class_names.get(cid, cid),
                round(actual["weekday"], 2),
                round(plan["weekday"], 2),
                round(actual["weekday"] - plan["weekday"], 2),
                round(actual["saturday"], 2),
                round(plan["saturday"], 2),
                round(actual["saturday"] - plan["saturday"], 2),
                round(actual["evening"], 2),
                round(plan["evening"], 2),
                round(actual["evening"] - plan["evening"], 2),
            ]
            for c, v in enumerate(vals, 1):
                cell = rel.cell(row=i, column=c, value=v)
                cell.font = cell_font
                cell.alignment = center
                cell.border = thin
                if c in {6, 9, 12} and isinstance(v, (int, float)) and abs(v) > 0.01:
                    cell.fill = diff_fill
        for idx, width in enumerate((10, 8, 14, 10, 10, 8, 10, 10, 8, 10, 10, 8), start=1):
            rel.column_dimensions[get_column_letter(idx)].width = width

    hours_rows: list[dict[str, Any]] = []
    if include & {"teacher_hours", "teacher_grid"}:
        hours_stats: dict[tuple[int, int], dict[str, Any]] = {}
        for slot in slots:
            if slot.teacher_id is None or slot.subject_id in activity_subject_ids:
                continue
            key = (slot.subject_id, slot.teacher_id)
            bucket = hours_stats.setdefault(
                key, {"classes": set(), "weekday": 0.0, "saturday": 0.0, "evening": 0.0}
            )
            bucket["classes"].add(slot.class_id)
            w = _wt(slot.week_parity)
            if slot.period >= evening_start:
                bucket["evening"] += w
            elif slot.weekday == 6:
                bucket["saturday"] += w
            else:
                bucket["weekday"] += w
        order_index = {name: i for i, name in enumerate(SUBJECT_ORDER)}
        for (sid, tid), st in hours_stats.items():
            sname = subject_names.get(sid, str(sid))
            hours_rows.append(
                {
                    "subject": sname,
                    "subject_order": order_index.get(sname, 999),
                    "teacher": teacher_names.get(tid, str(tid)),
                    "classes": "、".join(
                        sorted(
                            (_class_no(class_names[c]) for c in st["classes"] if c in class_names),
                            key=lambda x: int(x) if str(x).isdigit() else 0,
                        )
                    ),
                    "weekday": round(st["weekday"], 2),
                    "saturday": round(st["saturday"], 2),
                    "evening": round(st["evening"], 2),
                }
            )
        hours_rows.sort(key=lambda r: (r["subject_order"], r["teacher"]))

    if "teacher_hours" in include:
        hours = wb.create_sheet("教师课时")
        headers = ["学科", "序号", "姓名", "所教班级", "周一到周五课时", "周六课时", "周一到周六晚课课时"]
        for c, h in enumerate(headers, 1):
            cell = hours.cell(row=1, column=c, value=h)
            cell.fill = head_fill
            cell.font = bold
            cell.alignment = center
            cell.border = thin
        hours.auto_filter.ref = f"A1:G{1 + len(hours_rows)}"
        subject_spans: list[tuple[str, int, int]] = []
        cur_subject = None
        start = 2
        for i, r in enumerate(hours_rows):
            row = i + 2
            if cur_subject is None:
                cur_subject = r["subject"]
                start = row
            elif r["subject"] != cur_subject:
                subject_spans.append((cur_subject, start, row - 1))
                cur_subject = r["subject"]
                start = row
            vals = [
                r["subject"],
                i + 1,
                r["teacher"],
                r["classes"],
                int(r["weekday"]) if r["weekday"] == int(r["weekday"]) else r["weekday"],
                int(r["saturday"]) if r["saturday"] == int(r["saturday"]) else r["saturday"],
                int(r["evening"]) if r["evening"] == int(r["evening"]) else r["evening"],
            ]
            for c, v in enumerate(vals, 1):
                cell = hours.cell(row=row, column=c, value=v)
                cell.font = normal
                cell.alignment = center
                cell.border = thin
        if cur_subject is not None and hours_rows:
            subject_spans.append((cur_subject, start, 1 + len(hours_rows)))
        for _name, a, b in subject_spans:
            if a < b:
                hours.merge_cells(start_row=a, start_column=1, end_row=b, end_column=1)
                hours.cell(row=a, column=1).alignment = center
        for i, w in enumerate((8, 6, 12, 14, 14, 10, 16), start=1):
            hours.column_dimensions[get_column_letter(i)].width = w

    if "teacher_grid" in include:
        grid = wb.create_sheet("教师分课表")
        hours_ts: dict[tuple[int, int], float] = defaultdict(float)
        for slot in slots:
            if slot.teacher_id is None or slot.subject_id in activity_subject_ids:
                continue
            hours_ts[(slot.teacher_id, slot.subject_id)] += _wt(slot.week_parity)
        main_subject: dict[int, int] = {}
        for (tid, sid), h in hours_ts.items():
            cur = main_subject.get(tid)
            if cur is None or h > hours_ts[(tid, cur)]:
                main_subject[tid] = sid
        by_block: dict[tuple[int, int], list[PackSlot]] = defaultdict(list)
        for slot in slots:
            if slot.teacher_id is None:
                continue
            sid = slot.subject_id
            if sid in activity_subject_ids:
                sid = main_subject.get(slot.teacher_id) or sid
                if sid in activity_subject_ids:
                    continue
            by_block[(slot.teacher_id, sid)].append(slot)
        subj_rank = {n: i for i, n in enumerate(SUBJECT_ORDER)}
        block_keys = sorted(
            by_block.keys(),
            key=lambda k: (
                subj_rank.get(subject_names.get(k[1], ""), 99),
                teacher_names.get(k[0], ""),
            ),
        )
        cols_per_block = 8
        gap = 1
        blocks_per_row = 3
        day_periods = max(1, evening_start - 1)
        rows_per_block = 2 + day_periods + 2
        row_gap = 1
        total_cols = blocks_per_row * cols_per_block + (blocks_per_row - 1) * gap
        grid.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(total_cols, 1))
        grid_title = pack_title.replace("完整包", "教师分课表")
        if "教师分课表" not in grid_title:
            grid_title = "教师分课表"
        grid["A1"] = grid_title
        grid["A1"].font = Font(name="微软雅黑", bold=True, size=16)
        grid["A1"].alignment = center

        def write_teacher_block(tid: int, sid: int, start_row: int, start_col: int) -> None:
            tname = teacher_names.get(tid, f"教师{tid}")
            sname = subject_names.get(sid, "?")
            fill_hex = SUBJECT_HEADER_FILL.get(sname, "4472C4")
            block_fill = PatternFill("solid", fgColor=fill_hex)
            grid.merge_cells(
                start_row=start_row,
                start_column=start_col,
                end_row=start_row,
                end_column=start_col + cols_per_block - 1,
            )
            title = grid.cell(row=start_row, column=start_col, value=f"{sname} {tname}")
            title.font = white_bold
            title.fill = block_fill
            title.alignment = center
            for c in range(start_col, start_col + cols_per_block):
                cell = grid.cell(row=start_row, column=c)
                cell.fill = block_fill
                cell.border = thin
            hr = start_row + 1
            lab = grid.cell(row=hr, column=start_col, value="节/周")
            lab.font = bold
            lab.fill = head_fill
            lab.alignment = center
            lab.border = thin
            for i, name in enumerate(DAY_HEADERS):
                cell = grid.cell(row=hr, column=start_col + 1 + i, value=name)
                cell.font = bold
                cell.fill = head_fill
                cell.alignment = center
                cell.border = thin
            day_map: dict[tuple[int, int], list[tuple[str, str]]] = defaultdict(list)
            odd_eve: dict[int, list[tuple[str, str]]] = defaultdict(list)
            even_eve: dict[int, list[tuple[str, str]]] = defaultdict(list)
            for it in by_block[(tid, sid)]:
                cname = _class_no(class_names.get(it.class_id, str(it.class_id)))
                parity = it.week_parity
                entry = (cname, parity)
                if it.period >= evening_start:
                    if parity == "all":
                        odd_eve[it.weekday].append(entry)
                        even_eve[it.weekday].append(entry)
                    elif parity == "odd":
                        odd_eve[it.weekday].append(entry)
                    else:
                        even_eve[it.weekday].append(entry)
                    continue
                if it.weekday <= 5:
                    day_map[(it.weekday, it.period)].append(entry)
                elif it.weekday == 6:
                    if parity == "all":
                        day_map[(61, it.period)].append(entry)
                        day_map[(62, it.period)].append(entry)
                    elif parity == "odd":
                        day_map[(61, it.period)].append(entry)
                    else:
                        day_map[(62, it.period)].append(entry)
            for period in range(1, day_periods + 1):
                r = start_row + 1 + period
                lab = grid.cell(row=r, column=start_col, value=period)
                lab.font = bold
                lab.alignment = center
                lab.border = thin
                for wd in range(1, 6):
                    text = _cell_label(day_map.get((wd, period), []))
                    cell = grid.cell(row=r, column=start_col + wd, value=text or None)
                    cell.font = cell_font
                    cell.alignment = center
                    cell.border = thin
                for offset, key in ((6, 61), (7, 62)):
                    text = _cell_label(day_map.get((key, period), []))
                    cell = grid.cell(row=r, column=start_col + offset, value=text or None)
                    cell.font = cell_font
                    cell.alignment = center
                    cell.border = thin
            for label, bucket, row_offset, is_odd in (
                (evening_start, odd_eve, day_periods + 2, True),
                (f"{evening_start}双", even_eve, day_periods + 3, False),
            ):
                r = start_row + row_offset
                lab = grid.cell(row=r, column=start_col, value=label)
                lab.font = bold
                lab.fill = eve_fill
                lab.alignment = center
                lab.border = thin
                for wd in range(1, 6):
                    text = _cell_label(bucket.get(wd, []))
                    cell = grid.cell(row=r, column=start_col + wd, value=text or None)
                    cell.font = cell_font
                    cell.fill = eve_fill
                    cell.alignment = center
                    cell.border = thin
                for offset in (6, 7):
                    cell = grid.cell(row=r, column=start_col + offset, value=None)
                    cell.fill = eve_fill
                    cell.border = thin
                sat_text = _cell_label(bucket.get(6, []))
                sat_col = start_col + (6 if is_odd else 7)
                grid.cell(row=r, column=sat_col, value=sat_text or None).font = cell_font
                grid.cell(row=r, column=sat_col).fill = eve_fill
                grid.cell(row=r, column=sat_col).alignment = center

        for idx, (tid, sid) in enumerate(block_keys):
            block_row = idx // blocks_per_row
            block_col = idx % blocks_per_row
            start_row = 3 + block_row * (rows_per_block + row_gap)
            start_col = 1 + block_col * (cols_per_block + gap)
            write_teacher_block(tid, sid, start_row, start_col)
        for col in range(1, total_cols + 1):
            grid.column_dimensions[get_column_letter(col)].width = 5

    if "class_timetables" in include:
        for cid, cname in class_rows:
            title = _safe_sheet_name(wb, _sheet_title(cname))
            ws = wb.create_sheet(title)
            ws.merge_cells("A1:H1")
            ws["A1"] = f"班级：{cname}"
            ws["A1"].font = bold
            ws["A1"].fill = title_fill
            ws["A1"].alignment = center
            ws["A2"] = "节\\周"
            ws["A2"].font = bold
            ws["A2"].fill = head_fill
            ws["A2"].alignment = center
            ws["A2"].border = thin
            for col, name in enumerate(DAY_HEADERS, start=2):
                cell = ws.cell(row=2, column=col, value=name)
                cell.font = bold
                cell.fill = head_fill
                cell.alignment = center
                cell.border = thin
            day_map: dict[tuple[int, int], list[tuple[str, str]]] = defaultdict(list)
            odd_eve: dict[int, list[tuple[str, str]]] = defaultdict(list)
            even_eve: dict[int, list[tuple[str, str]]] = defaultdict(list)
            for it in slots:
                if it.class_id != cid:
                    continue
                sname = it.subject_name
                if it.subject_id in activity_subject_ids:
                    sname = "班主任晚课" if it.teacher_id is not None else "自主学习"
                short = _short_subject(sname)
                parity = it.week_parity
                if it.period >= evening_start:
                    if parity == "all":
                        odd_eve[it.weekday].append((short, parity))
                        even_eve[it.weekday].append((short, parity))
                    elif parity == "odd":
                        odd_eve[it.weekday].append((short, parity))
                    else:
                        even_eve[it.weekday].append((short, parity))
                    continue
                if it.weekday <= 5:
                    day_map[(it.weekday, it.period)].append((short, parity))
                elif it.weekday == 6:
                    if parity == "all":
                        day_map[(61, it.period)].append((short, parity))
                        day_map[(62, it.period)].append((short, parity))
                    elif parity == "odd":
                        day_map[(61, it.period)].append((short, parity))
                    else:
                        day_map[(62, it.period)].append((short, parity))
            max_day_period = max(1, evening_start - 1)
            for period in range(1, max_day_period + 1):
                r = period + 2
                lab = ws.cell(row=r, column=1, value=period)
                lab.font = bold
                lab.alignment = center
                lab.border = thin
                for wd in range(1, 6):
                    text = _cell_label(day_map.get((wd, period), []))
                    cell = ws.cell(row=r, column=wd + 1, value=text or None)
                    cell.font = pe_font if text == "体" else normal
                    cell.alignment = center
                    cell.border = thin
                for col, key in ((7, 61), (8, 62)):
                    text = _cell_label(day_map.get((key, period), []))
                    cell = ws.cell(row=r, column=col, value=text or None)
                    cell.font = pe_font if text == "体" else normal
                    cell.alignment = center
                    cell.border = thin
            eve_base = max_day_period + 2
            for label, bucket, row_offset, is_odd in (
                (evening_start, odd_eve, 1, True),
                (f"{evening_start}双", even_eve, 2, False),
            ):
                r = eve_base + row_offset
                lab = ws.cell(row=r, column=1, value=label)
                lab.font = bold
                lab.fill = eve_fill
                lab.alignment = center
                lab.border = thin
                for wd in range(1, 6):
                    text = _cell_label(bucket.get(wd, []))
                    cell = ws.cell(row=r, column=wd + 1, value=text or None)
                    cell.font = normal
                    cell.fill = eve_fill
                    cell.alignment = center
                    cell.border = thin
                for col in (7, 8):
                    cell = ws.cell(row=r, column=col, value=None)
                    cell.fill = eve_fill
                    cell.border = thin
                    cell.alignment = center
                sat_text = _cell_label(bucket.get(6, []))
                sat_col = 7 if is_odd else 8
                ws.cell(row=r, column=sat_col, value=sat_text or None).font = normal
                ws.cell(row=r, column=sat_col).fill = eve_fill
                ws.cell(row=r, column=sat_col).alignment = center
            ws.column_dimensions["A"].width = 8
            for col in range(2, 9):
                ws.column_dimensions[get_column_letter(col)].width = 6
            for r in range(1, eve_base + 3):
                ws.row_dimensions[r].height = 22

    if not wb.sheetnames:
        raise ValueError("请至少勾选一种工作表")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
